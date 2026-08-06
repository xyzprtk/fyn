import csv
import io
import sys
import time
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "parser"))

from extractors import extract  # noqa: E402


def csv_bytes(row_count: int = 5000) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Date", "Description", "Debit", "Credit", "Balance"])
    for index in range(row_count):
        writer.writerow(["01/06/2026", f"PERF MERCHANT {index}", "125.00", "", "100000.00"])
    return output.getvalue().encode()


def xlsx_bytes(row_count: int = 2000) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["Date", "Description", "Debit", "Credit", "Balance"])
    for index in range(row_count):
        sheet.append(["01/06/2026", f"PERF MERCHANT {index}", "125.00", "", "100000.00"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(page_count: int = 3, rows_per_page: int = 40) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_font("helvetica", size=7)
    for page_number in range(page_count):
        pdf.add_page()
        for row_number in range(rows_per_page):
            cells = [
                "01/06/2026",
                f"PERF MERCHANT {page_number * rows_per_page + row_number}",
                "125.00",
                "",
                "100000.00",
            ]
            for cell in cells:
                pdf.cell(35, 4, cell, border=1)
            pdf.ln()
    return bytes(pdf.output())


def measure(kind: str, data: bytes, filename: str) -> None:
    started = time.perf_counter()
    raw = extract(kind, data, filename)
    duration = (time.perf_counter() - started) * 1000
    print(f"{kind}\tpages={raw.page_count}\trows={len(raw.rows)}\tbytes={len(data)}\tms={duration:.2f}")


print("kind\tmetadata")
measure("csv", csv_bytes(), "performance.csv")
measure("xlsx", xlsx_bytes(), "performance.xlsx")
measure("pdf", pdf_bytes(), "performance.pdf")
