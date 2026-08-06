"""CSV and XLSX extraction.

CSV uses the stdlib csv module on purpose: bank CSV exports are frequently
ragged (preamble lines with fewer columns than the table), which pandas'
C engine rejects. XLSX uses openpyxl's read-only iterator so large workbooks do
not need to be materialized as a pandas DataFrame.
"""

import csv
import io

from extractors import ExtractionError
from models import RawDoc


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_csv(data: bytes, filename: str) -> RawDoc:
    text = _decode(data)
    if not text.strip():
        raise ExtractionError("file is empty")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = [
        [cell.strip() for cell in row]
        for row in csv.reader(io.StringIO(text), dialect)
    ]
    if not rows:
        raise ExtractionError("no rows found in csv")
    return RawDoc(filename=filename, kind="csv", rows=rows, text=text, page_count=1)


def extract_xlsx(data: bytes, filename: str) -> RawDoc:
    workbook = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    except Exception as exc:
        raise ExtractionError(f"could not read xlsx: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()

    if not rows:
        raise ExtractionError("no rows found in xlsx")
    return RawDoc(filename=filename, kind="xlsx", rows=rows, text="", page_count=1)
